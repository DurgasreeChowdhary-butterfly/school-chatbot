"""Three fake, messy school documents for testing scripts/ingest_docs.py (all DUMMY content):
- fee-circular-2027.pdf: a fee circular (PDF), with a tentative amount and a sentence the PDF wraps
- calendar-2027-28.docx: an academic calendar (Word) with a holiday table and an open-ended vacation
- notice-board-photo.txt: text as it comes out of a phone photo of a notice board (OCR mistakes, a personal
  mobile number, a relative date) and a fee that CONFLICTS with the circular (Rs 27,000 vs Rs 26,000)

Make them by hand:  python -m tests.messy_docs incoming"""
import sys
from pathlib import Path

FEE_LINES = [
    "DEMO INTERNATIONAL SCHOOL, CAMPUS A",
    "FEE CIRCULAR 2027-28",
    "Dear Parents, the fee structure for the academic year 2027-28 is given below.",
    "Nursery to UKG: Admission fee Rs 10,000; Annual fee Rs 21,000; Term fee Rs 15,500",
    "Grade 1 to 5: Admission fee Rs 12,000; Annual fee Rs 26,000; Term fee Rs 18,500",
    "Grade 6 to 8: Admission fee Rs 15,000; Annual fee Rs 31,000; Term fee Rs 22,500",
    "Grade 9 to 10: Admission fee Rs 18,000; Annual fee Rs 36,000; Term fee Rs 26,500",
    "Term fees are due on 10 June 2027, 10 September 2027 and 10 December 2027.",
    "Fees may be paid online through the school portal or at the accounts office between 9.00 am and 3.00 pm on "
    "working days, and parents are requested to keep the receipt safely for future reference.",
    "A late fee of Rs 100 per day will apply (amount tentative).",
    "Sibling concession: 5% of the annual fee for the second child, with the principal's written approval.",
    "Duplicate fee receipts are issued by the accounts office on request.",
    "For fee queries write to accounts.campusa@example.com",
]

CALENDAR_PARAGRAPHS = [
    "The academic year begins on 1 June 2027.",
    "Term 1 examinations: 13 to 20 September 2027",
    "Term 2 examinations: 9 to 17 December 2027",
    "Annual examinations: March 2028 (dates to be announced)",
    "Parent-teacher meetings are held on the last Saturday of September, December and March, 9:30 am to 12:30 pm.",
    "Working Saturdays: first and third Saturday of every month.",
]
HOLIDAYS = [("Holiday", "From", "To"),
            ("Dussehra", "18 October 2027", "20 October 2027"),
            ("Diwali", "27 October 2027", "30 October 2027"),
            ("Christmas", "25 December 2027", "25 December 2027"),
            ("Sankranti", "13 January 2028", "16 January 2028"),
            ("Summer vacation", "from 1 May 2028", "")]

NOTICE = """DEM0 INTERNATI0NAL SCH00L - N0TICE B0ARD
Office hours: Monday to Friday 8:30 am to 4:30 pm, Saturday 8:30 am to 1:00 pm
Front office phone: +91 00000 00001
Admission counsellor: Ms. Sample Counsellor, admissions.campusa@example.com
Annual fee Grade 1-5: Rs 27,000 (revised)
Bus service within 12 km. New stop requests to the transport coordinator, transport.campusa@example.com
Schoo1 timings: 8:45 am to 3:15 pm (Grade 1 and above); pre-primary til1 12:30
Hostel: n0t available
PTM next Saturday ~~ details fr0m class teacher ??
For uniform queries contact Mrs Sample Teacher 9876501234
"""


def make_messy_docs(folder: Path) -> list[Path]:
    from docx import Document
    from fpdf import FPDF

    folder.mkdir(parents=True, exist_ok=True)
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in FEE_LINES:
        pdf.multi_cell(0, 7, line, new_x="LMARGIN", new_y="NEXT")  # long lines wrap, like a real circular
    pdf_path = folder / "fee-circular-2027.pdf"
    pdf.output(str(pdf_path))

    doc = Document()
    doc.add_heading("ACADEMIC CALENDAR 2027-28", level=1)
    for p in CALENDAR_PARAGRAPHS:
        doc.add_paragraph(p)
    table = doc.add_table(rows=0, cols=3)
    for row in HOLIDAYS:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    docx_path = folder / "calendar-2027-28.docx"
    doc.save(docx_path)

    txt_path = folder / "notice-board-photo.txt"
    txt_path.write_text(NOTICE, encoding="utf-8")
    return [pdf_path, docx_path, txt_path]


if __name__ == "__main__":
    for p in make_messy_docs(Path(sys.argv[1] if len(sys.argv) > 1 else "incoming")):
        print("made", p)
